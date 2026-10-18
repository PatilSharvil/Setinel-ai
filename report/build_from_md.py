from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)

# title
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run('AST — AI Security & Code Assistant')
r.bold = True; r.font.size = Pt(22); r.font.color.rgb = RGBColor(0x1A,0x73,0xE8)
s = doc.add_paragraph(); s.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s.add_run('Project Report (Detailed)'); r.font.size = Pt(14)
doc.add_page_break()

md = open('project.md', encoding='utf-8').read()
lines = md.splitlines()
for line in lines:
    line = line.rstrip()
    if line.startswith('# '):
        doc.add_page_break()
        h = doc.add_heading(line[2:], level=1)
        for run in h.runs: run.font.name = 'Times New Roman'
    elif line.startswith('## '):
        doc.add_page_break()
        h = doc.add_heading(line[3:], level=2)
        for run in h.runs: run.font.name = 'Times New Roman'
    elif line.startswith('### '):
        h = doc.add_heading(line[4:], level=3)
        for run in h.runs: run.font.name = 'Times New Roman'
    elif line.strip() == '':
        doc.add_paragraph('')
    else:
        clean = ''.join(c for c in line if c == '\t' or c == '\n' or (ord(c) >= 32 and ord(c) != 127))
        p = doc.add_paragraph(clean)
        for run in p.runs:
            run.font.name = 'Times New Roman'

# Insert diagrams at end of architecture and ReAct sections? Add appendix images
doc.add_page_break()
doc.add_heading('Appendix: Architecture Diagram', level=1)
doc.add_picture('report/fig/arch.png', width=Inches(6.2))
cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = cap.add_run('Figure 1: High-level architecture of AST'); r.italic=True; r.font.size=Pt(10)

doc.add_page_break()
doc.add_heading('Appendix: ReAct Loop Diagram', level=1)
doc.add_picture('report/fig/react.png', width=Inches(5.2))
cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = cap.add_run('Figure 2: ReAct graph structure'); r.italic=True; r.font.size=Pt(10)

doc.save('report/AST_Detailed_Report.docx')
print('saved')
