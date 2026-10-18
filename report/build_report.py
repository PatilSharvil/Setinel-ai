from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)

def add_heading(text, level=1):
    if level == 1:
        doc.add_page_break()
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = 'Times New Roman'
    return h

def add_para(text, bold=False, italic=False, size=12, align=None):
    p = doc.add_paragraph()
    if align: p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    return p

def add_bullet(text):
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(text)
    return p

def add_image(path, caption, width=6.0):
    doc.add_picture(path, width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(10)
    return cap

# Title page
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run('AST — AI Security & Code Assistant')
r.bold = True
r.font.size = Pt(24)
r.font.color.rgb = RGBColor(0x1A,0x73,0xE8)
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run('Project Report')
r.font.size = Pt(16)
sub2 = doc.add_paragraph()
sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub2.add_run('ReAct-Based Terminal AI Coding and Security Assistant')
r.italic = True
r.font.size = Pt(12)
add_para('Author: AST Development Team', size=11)
add_para('Version: 1.0 (MVP)', size=11)
doc.add_page_break()

# Abstract
add_heading('Abstract', 1)
add_para('AST is a terminal-based AI coding and security assistant built with a genuine ReAct '
         '(Reasoning + Acting) loop using LangGraph and LangChain. It operates on the current '
         'project directory, discovers files, runs security scanners (Gitleaks, Semgrep, ripgrep), '
         'and provides natural-language explanations and remediation guidance. The system is '
         'read-only by default, enforces path safety and timeouts, and supports Groq and OpenRouter '
         'providers with automatic fallback. This report describes the architecture, tools, '
         'security boundaries, installation, testing, and limitations of the MVP.')

# Table of contents
add_heading('Table of Contents', 1)
toc = [
    '1. Introduction',
    '2. Objectives',
    '3. System Architecture',
    '4. ReAct Agent Loop',
    '5. Tools',
    '6. Security Boundaries',
    '7. LLM Providers',
    '8. Installation and Usage',
    '9. Testing',
    '10. Sample Findings',
    '11. Limitations',
    '12. Conclusion',
]
for item in toc:
    add_para(item)

# 1 Introduction
add_heading('1. Introduction', 1)
add_para('AST (AI Security & Code Assistant) is a minimal MVP of an interactive, project-aware '
         'coding assistant specialized in understanding codebases and performing security audits. '
         'Unlike fixed command-line scanners, AST uses a LangGraph ReAct agent that decides which '
         'tools to invoke based on the user request and prior observations. The assistant runs in '
         'any project directory via a global `ast` command and keeps conversation context for '
         'follow-up questions.')

# 2 Objectives
add_heading('2. Objectives', 1)
add_bullet('Provide a global `ast` command usable in any project directory on Windows.')
add_bullet('Implement a genuine ReAct tool-calling loop with LangGraph and LangChain.')
add_bullet('Integrate Gitleaks, Semgrep, and ripgrep as safe subprocess tools.')
add_bullet('Preserve conversation context and scan findings across follow-up questions.')
add_bullet('Support Groq and OpenRouter providers with fallback and clear error reporting.')
add_bullet('Enforce read-only operation, path safety, timeouts, and iteration limits.')
add_bullet('Render Markdown responses in the terminal with Rich.')

# 3 Architecture
add_heading('3. System Architecture', 1)
add_para('The application consists of a Typer CLI, a Rich terminal UI, a LangGraph ReAct agent, '
         'a set of validated tools, a findings store, and provider adapters. The following diagram '
         'shows the high-level architecture.')
add_image('report/fig/arch.png', 'Figure 1: High-level architecture of AST', width=6.2)

# 4 ReAct
add_heading('4. ReAct Agent Loop', 1)
add_para('The agent follows a genuine ReAct cycle: the LLM receives the conversation and tool '
         'definitions, decides whether to call a tool, the tool executes and returns an '
         'observation, and the LLM continues until it produces a final answer. The graph uses a '
         'conditional edge between the agent node and the tool node. A bounded recursion limit '
         'prevents infinite loops.')
add_image('report/fig/react.png', 'Figure 2: ReAct graph structure', width=5.2)

# 5 Tools
add_heading('5. Tools', 1)
tools_table = doc.add_table(rows=1, cols=3)
tools_table.style = 'Table Grid'
hdr = tools_table.rows[0].cells
hdr[0].text = 'Tool'
hdr[1].text = 'Purpose'
hdr[2].text = 'Key Safety Controls'
rows = [
    ('list_files', 'Understand project structure', 'Excludes build dirs, no symlinks, limit 300'),
    ('read_file', 'Inspect source code', 'Path containment, line/byte limits, secret redaction'),
    ('search_code', 'Fast code search via ripgrep', 'Arg list (no shell), result cap, timeout'),
    ('scan_secrets', 'Run Gitleaks', 'JSON parsing, redaction, git-tracked annotation'),
    ('scan_vulnerabilities', 'Run Semgrep', 'Timeout handling, malformed output handling'),
    ('search_security_patterns', 'Heuristic security patterns', 'Focused regex set, excluded dirs'),
    ('get_findings', 'Retrieve stored findings', 'Filters by severity/category/source/id'),
]
for r in rows:
    cells = tools_table.add_row().cells
    cells[0].text = r[0]
    cells[1].text = r[1]
    cells[2].text = r[2]
add_para('Table 1: AST tools and safety controls', italic=True, size=10)

# 6 Security
add_heading('6. Security Boundaries', 1)
add_bullet('Read-only by default; no automatic file modification.')
add_bullet('No arbitrary shell execution; only known scanner binaries via argument lists.')
add_bullet('Path traversal and external symlink access are blocked.')
add_bullet('Secrets are redacted before display or model transmission.')
add_bullet('Bounded iterations (default 10 rounds) and per-tool timeouts.')
add_bullet('Provider API keys are never logged.')
add_bullet('Git-tracked credential findings are annotated as compromised.')

# 7 Providers
add_heading('7. LLM Providers', 1)
add_para('Groq is supported via langchain_groq.ChatGroq and OpenRouter via '
         'langchain_openai.ChatOpenAI with the OpenRouter base URL. The provider is selected '
         'through environment variables, with an automatic fallback to the secondary provider '
         'on failure. Models are configurable; `openrouter/free` is supported.')
prov = doc.add_table(rows=1, cols=3)
prov.style = 'Table Grid'
h = prov.rows[0].cells
h[0].text = 'Setting'
h[1].text = 'Groq'
h[2].text = 'OpenRouter'
prows = [
    ('Env var', 'GROQ_API_KEY', 'OPENROUTER_API_KEY'),
    ('Model example', 'llama-3.1-8b-instant', 'openrouter/free'),
    ('Base URL', 'Groq default', 'https://openrouter.ai/api/v1'),
]
for r in prows:
    cells = prov.add_row().cells
    cells[0].text = r[0]
    cells[1].text = r[1]
    cells[2].text = r[2]
add_para('Table 2: Provider configuration', italic=True, size=10)

# 8 Installation
add_heading('8. Installation and Usage', 1)
add_para('Install into a virtual environment and refresh the shell:')
code_lines = [
    'cd "C:\\projects\\appsentinel v2"',
    'python -m venv .venv',
    '.\\.venv\\Scripts\\Activate.ps1',
    'python -m pip install -e .',
    'cd C:\\projects\\my-app',
    'ast',
]
for line in code_lines:
    p = doc.add_paragraph()
    r = p.add_run(line)
    r.font.name = 'Consolas'
    r.font.size = Pt(10)
add_para('Slash commands: /scan, /findings, /clear, /provider, /help, /exit. Natural language '
         'remains the primary interaction method.')

# 9 Testing
add_heading('9. Testing', 1)
add_para('Unit tests cover findings storage, tool path safety and redaction, and the agent loop '
         'with a mocked LLM. End-to-end tests verified the global `ast` command, OpenRouter '
         'tool calling, Gitleaks parsing, ripgrep search, and Markdown rendering on Windows.')
test_rows = [
    ('test_findings', 'Add/filter findings, summary', 'Pass'),
    ('test_tools', 'Path traversal, redaction, list_files', 'Pass'),
    ('test_agent', 'Agent final answer without tool', 'Pass'),
    ('E2E scan', 'Gitleaks + pattern scan on sample project', 'Pass'),
]
t = doc.add_table(rows=1, cols=3)
t.style = 'Table Grid'
h = t.rows[0].cells
h[0].text = 'Test'
h[1].text = 'Coverage'
h[2].text = 'Result'
for r in test_rows:
    cells = t.add_row().cells
    cells[0].text = r[0]
    cells[1].text = r[1]
    cells[2].text = r[2]
add_para('Table 3: Test results', italic=True, size=10)

# 10 Sample findings
add_heading('10. Sample Findings', 1)
add_para('A typical scan produces findings such as an exposed API key in a Git-tracked .env '
         'file, hardcoded secrets, and dangerous eval/exec patterns. Each finding includes an '
         'ID, severity, path, line number, and redacted evidence. Findings persist in memory for '
         'follow-up questions like "Explain finding F003".')
add_bullet('F001 [High] Generic API Key — backend/.env:27 (git-tracked: yes)')
add_bullet('F011-F017 [Low] eval() usage — potential code execution')
add_bullet('F018-F024 [Low] exec() usage — potential command injection')

# 11 Limitations
add_heading('11. Limitations', 1)
add_bullet('Semgrep is optional and may not be installed; the tool reports this accurately.')
add_bullet('No persistent database; state is in-process only.')
add_bullet('Large repositories are truncated by file and output limits.')
add_bullet('Page numbers in this report were not verified without LibreOffice.')

# 12 Conclusion
add_heading('12. Conclusion', 1)
add_para('AST delivers a functional, safe, terminal-based ReAct security assistant with real '
         'tool execution, structured findings, and provider flexibility. It meets the MVP goals '
         'of project understanding and security auditing while maintaining strict execution '
         'boundaries.')

# Appendix A
add_heading('Appendix A: Module Descriptions', 1)
modules = [
    ('cli.py', 'Typer entry point, Rich UI, slash commands, spinner, Markdown rendering.'),
    ('config.py', 'Environment-based settings, project root resolution, excluded directories.'),
    ('providers.py', 'Groq and OpenRouter LLM construction with fallback support.'),
    ('tools.py', 'LangChain tools with path safety, redaction, and activity logging.'),
    ('scanner.py', 'Gitleaks, Semgrep, and ripgrep execution and parsing.'),
    ('findings.py', 'Finding dataclass and in-memory FindingsStore.'),
    ('agent.py', 'LangGraph ReAct graph, Session state, bounded execution.'),
]
for name, desc in modules:
    add_para(name + ': ' + desc, bold=False)
    add_para('')

# Appendix B
add_heading('Appendix B: Sample Interaction', 1)
transcript = [
    'You > Scan this project for security vulnerabilities.',
    'AST > I\'ll inspect the project and run the relevant security tools.',
    '     [1] Discovering project files...',
    '     [2] Running Gitleaks...',
    '     [3] Running Semgrep...',
    '     [4] Searching suspicious code patterns...',
    'AST > The scan completed. Here are the findings...',
    'You > Explain the highest-risk finding.',
    'AST > I\'ll inspect the affected code and explain the evidence.',
    'You > Find other places where the same problem occurs.',
    'AST > I\'ll search the codebase and compare the results.',
    'You > What should I fix first?',
    'AST > Based on the findings, I recommend addressing...',
]
for line in transcript:
    p = doc.add_paragraph()
    r = p.add_run(line)
    r.font.name = 'Consolas'
    r.font.size = Pt(10)

# Appendix C
add_heading('Appendix C: Security Checklist', 1)
for item in [
    'No automatic file modification.',
    'No arbitrary shell execution.',
    'Path traversal blocked.',
    'Symlinks not followed.',
    'Secrets redacted in observations.',
    'API keys never logged.',
    'Timeouts enforced on scanners.',
    'Iteration limit enforced.',
    'Missing tools reported accurately.',
    'Provider failures reported with fallback notice.',
]:
    add_bullet(item)

# Appendix D
add_heading('Appendix D: Error Handling', 1)
add_para('Provider errors, missing tools, timeouts, and malformed scanner output are caught and '
         'reported to the user without crashing the session. The agent stops gracefully when the '
         'recursion limit is reached. Fallback providers are attempted only once and announced '
         'explicitly to avoid silent switching.')
add_para('All subprocess calls use argument lists and UTF-8 decoding with replacement to avoid '
         'encoding crashes on Windows. Tool outputs are truncated to 8000 characters to protect '
         'the model context window.')

doc.save('report/AST_Project_Report_v2.docx')
print('saved')

print('saved')
